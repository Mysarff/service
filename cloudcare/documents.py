"""Validated document loading, CPU OCR and LangChain parent/child splitting.

The educational loaders supplied the original format coverage.  This module
keeps that coverage in the customer-support package and adds source hashes,
bounded parsing and deterministic identifiers.  It never imports the old app.
"""
from __future__ import annotations

import hashlib
import io
import re
import threading
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


class DocumentValidationError(ValueError):
    """The supplied file is unsupported, disguised, unsafe or has no text."""


@dataclass
class IngestedDocument:
    id: str
    sha256: str
    source: str
    title: str
    category: str
    version: str
    documents: list[Document]
    parents: list[dict[str, Any]]
    chunks: list[dict[str, Any]]
    metadata: dict[str, Any] = field(default_factory=dict)

    def document_record(self) -> dict[str, Any]:
        """Metadata row for MySQL; extracted content lives in parent/child rows."""
        return {
            "id": self.id,
            "sha256": self.sha256,
            "source": self.source,
            "title": self.title,
            "category": self.category,
            "version": self.version,
            "content_chars": sum(len(d.page_content) for d in self.documents),
            "chunk_count": len(self.chunks),
            "metadata": self.metadata,
        }


class DocumentProcessor:
    """Load real TXT/MD/PDF/DOCX/PPTX/images and create indexed child units."""

    EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".webp"})

    def __init__(self, settings: Any):
        self.settings = settings
        self.max_bytes = int(getattr(settings, "max_upload_bytes", 10_000_000))
        self.max_pages = int(getattr(settings, "max_document_pages", 500))
        self.max_ocr_images = int(getattr(settings, "max_ocr_images", 200))
        self.max_pixels = int(getattr(settings, "max_image_pixels", 25_000_000))
        parent_size = int(getattr(settings, "parent_chunk_size", 1200))
        child_size = int(getattr(settings, "child_chunk_size", 300))
        overlap = int(getattr(settings, "chunk_overlap", 50))
        if not (0 <= overlap < child_size <= parent_size):
            raise ValueError("分块配置必须满足 0 <= overlap < child <= parent")
        separators = ["\n\n", "\n", "。", "！", "？", "；", ";", " ", ""]
        self.parent_splitter = RecursiveCharacterTextSplitter(
            chunk_size=parent_size, chunk_overlap=overlap,
            separators=separators, add_start_index=True,
        )
        self.child_splitter = RecursiveCharacterTextSplitter(
            chunk_size=child_size, chunk_overlap=overlap,
            separators=separators, add_start_index=True,
        )
        self._ocr = None
        self._ocr_lock = threading.Lock()

    def ingest_path(
        self, path: str | Path, category: str = "uploaded", source_name: str | None = None,
        version: str = "1", tenant_id: str = "demo", visibility: str = "public",
        synthetic: bool = False,
    ) -> IngestedDocument:
        file_path = Path(path).resolve(strict=True)
        if not file_path.is_file():
            raise DocumentValidationError("请提供文件路径")
        extension = file_path.suffix.lower()
        if extension not in self.EXTENSIONS:
            raise DocumentValidationError("仅支持 TXT、MD、PDF、DOCX、PPTX、PNG、JPEG、WebP；旧 DOC/PPT 请先转换")
        size = file_path.stat().st_size
        if not 0 < size <= self.max_bytes:
            raise DocumentValidationError(f"文件大小必须介于 1 和 {self.max_bytes} 字节之间")
        with file_path.open("rb") as stream:
            payload = stream.read(self.max_bytes + 1)
        if len(payload) > self.max_bytes:
            raise DocumentValidationError("文件超过上传限制")
        self._validate_signature(extension, payload)
        file_hash = hashlib.sha256(payload).hexdigest()
        display_source = source_name or file_path.name
        if len(display_source) > 500 or "\x00" in display_source:
            raise DocumentValidationError("来源名称过长或包含非法字符")
        document_id = "DOC-" + file_hash[:24]
        title = Path(display_source).stem or document_id
        if len(title) > 255 or len(category) > 128 or len(str(version)) > 64:
            raise DocumentValidationError("标题、业务分类或版本名称超过元数据长度限制")
        if not tenant_id or len(tenant_id) > 96 or visibility not in {"public", "internal", "private"}:
            raise DocumentValidationError("租户或可见性元数据不合法")
        base_metadata = {
            "document_id": document_id, "source": display_source, "title": title,
            "category": category, "version": str(version), "tenant_id": tenant_id,
            "visibility": visibility, "synthetic": bool(synthetic),
            "file_sha256": file_hash, "file_bytes": len(payload), "format": extension.lstrip("."),
        }
        documents = self._load(extension, payload, base_metadata)
        documents = [d for d in documents if d.page_content.strip()]
        if not documents:
            raise DocumentValidationError("文件未提取出可检索文本；请检查文档或扫描质量")
        if sum(len(d.page_content) for d in documents) > self.max_bytes * 20:
            raise DocumentValidationError("提取文本超过允许大小")
        parents: list[dict[str, Any]] = []
        chunks: list[dict[str, Any]] = []
        for page_index, document in enumerate(documents):
            for parent_index, parent in enumerate(self.parent_splitter.split_documents([document])):
                parent_id = f"{document_id}-P{page_index:04d}-{parent_index:04d}"
                parent_row = self._record(parent, parent_id, document_id, "", "", title)
                parents.append(parent_row)
                for child_index, child in enumerate(self.child_splitter.split_documents([parent])):
                    child_id = f"{parent_id}-C{child_index:04d}"
                    # The child start index is relative to its parent; retain both.
                    child.metadata["parent_start_index"] = parent.metadata.get("start_index", 0)
                    chunks.append(self._record(
                        child, child_id, document_id, parent_id, parent.page_content, title,
                    ))
        metadata = dict(base_metadata)
        metadata.update({
            "page_count": len(documents), "parent_count": len(parents),
            "chunk_count": len(chunks), "ocr_pages": sum(bool(d.metadata.get("ocr")) for d in documents),
            "ocr_engine": "RapidOCR ONNX Runtime" if any(d.metadata.get("ocr") for d in documents) else None,
            "splitter": "langchain_text_splitters.RecursiveCharacterTextSplitter",
        })
        return IngestedDocument(
            document_id, file_hash, display_source, title, category, str(version),
            documents, parents, chunks, metadata,
        )

    @staticmethod
    def _record(document: Document, row_id: str, document_id: str, parent_id: str,
                parent_content: str, title: str) -> dict[str, Any]:
        content = document.page_content
        metadata = dict(document.metadata)
        return {
            "id": row_id, "document_id": document_id, "parent_id": parent_id,
            "title": title, "category": metadata["category"], "content": content,
            "tags": [], "source": metadata["source"], "version": metadata["version"],
            "synthetic": metadata["synthetic"], "tenant_id": metadata["tenant_id"],
            "visibility": metadata["visibility"], "parent_content": parent_content,
            "content_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "metadata": metadata,
        }

    def _validate_signature(self, extension: str, payload: bytes) -> None:
        if extension == ".pdf":
            if not payload.lstrip().startswith(b"%PDF-"):
                raise DocumentValidationError("扩展名为 PDF，但文件头不是 PDF")
        elif extension in {".docx", ".pptx"}:
            required = "word/document.xml" if extension == ".docx" else "ppt/presentation.xml"
            try:
                with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                    entries = archive.infolist()
                    if len(entries) > 10_000 or sum(i.file_size for i in entries) > self.max_bytes * 20:
                        raise DocumentValidationError("Office 压缩内容超过解析限制")
                    names = set(archive.namelist())
                    if "[Content_Types].xml" not in names or required not in names:
                        raise DocumentValidationError("Office 内容与文件扩展名不匹配")
                    if any(i.flag_bits & 1 for i in entries):
                        raise DocumentValidationError("不支持加密 Office 文件")
                    if any("vbaproject" in name.lower() for name in names):
                        raise DocumentValidationError("不支持包含宏的 Office 文件")
            except zipfile.BadZipFile as exc:
                raise DocumentValidationError("文件不是有效的 Office 文档") from exc
        elif extension in {".txt", ".md"}:
            if payload.startswith((b"%PDF-", b"PK\x03\x04", b"\x89PNG", b"\xff\xd8\xff")) or b"\x00" in payload:
                raise DocumentValidationError("文本文件包含二进制内容或扩展名伪装")
            try:
                payload.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise DocumentValidationError("TXT/Markdown 必须为 UTF-8 编码") from exc
        else:
            from PIL import Image, UnidentifiedImageError
            expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[extension]
            try:
                with Image.open(io.BytesIO(payload)) as image:
                    if image.format != expected:
                        raise DocumentValidationError("图片内容与文件扩展名不匹配")
                    if image.width * image.height > self.max_pixels:
                        raise DocumentValidationError("图片像素数超过解析限制")
                    image.verify()
            except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
                raise DocumentValidationError("图片无法解析或超过像素限制") from exc

    def _load(self, extension: str, payload: bytes, metadata: dict[str, Any]) -> list[Document]:
        if extension in {".txt", ".md"}:
            return [Document(page_content=payload.decode("utf-8-sig"), metadata={**metadata, "page": 1, "ocr": False})]
        if extension == ".pdf":
            return self._load_pdf(payload, metadata)
        if extension == ".docx":
            return self._load_docx(payload, metadata)
        if extension == ".pptx":
            return self._load_pptx(payload, metadata)
        text, confidence = self._ocr_bytes(payload)
        return [Document(page_content=text, metadata={**metadata, "page": 1, "ocr": True, "ocr_confidence": confidence})]

    def _ocr_bytes(self, payload: bytes) -> tuple[str, float | None]:
        import numpy as np
        from PIL import Image
        with Image.open(io.BytesIO(payload)) as image:
            if image.width * image.height > self.max_pixels:
                raise DocumentValidationError("OCR 图片像素数超过限制")
            image_array = np.asarray(image.convert("RGB"))
        return self._ocr_array(image_array)

    def _ocr_array(self, image_array: Any) -> tuple[str, float | None]:
        with self._ocr_lock:
            if self._ocr is None:
                from rapidocr_onnxruntime import RapidOCR
                self._ocr = RapidOCR(det_use_cuda=False, cls_use_cuda=False, rec_use_cuda=False)
            result, _elapsed = self._ocr(image_array)
        lines = [str(line[1]).strip() for line in result or [] if str(line[1]).strip()]
        scores = [float(line[2]) for line in result or [] if len(line) > 2]
        return "\n".join(lines), (sum(scores) / len(scores) if scores else None)

    def _load_pdf(self, payload: bytes, metadata: dict[str, Any]) -> list[Document]:
        import fitz
        import numpy as np
        documents = []
        with fitz.open(stream=payload, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise DocumentValidationError("不支持加密 PDF")
            if pdf.page_count > self.max_pages:
                raise DocumentValidationError("PDF 页数超过解析限制")
            ocr_images = 0
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text", sort=True).strip()
                used_ocr = False
                confidence = None
                # A scan may have a searchable header while its actual body is
                # an image. Do not treat the presence of header text as proof
                # that all page content has already been extracted.
                page_area = max(1, page.rect.width * page.rect.height)
                dominant_image = any(
                    rect.width * rect.height / page_area >= 0.4
                    for image in page.get_images(full=True)
                    for rect in page.get_image_rects(image[0])
                )
                little_text = len(re.sub(r"\s", "", text)) < 20
                if little_text or dominant_image:
                    # Render the entire page: this also handles rotated scans and
                    # pages assembled from multiple small image tiles.
                    scale = min(2.0, (self.max_pixels / max(1, page.rect.width * page.rect.height)) ** 0.5)
                    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, colorspace=fitz.csRGB)
                    array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
                    ocr_images += 1
                    if ocr_images > self.max_ocr_images:
                        raise DocumentValidationError("扫描 PDF 超过 OCR 页数限制")
                    ocr_text, confidence = self._ocr_array(array)
                    if ocr_text:
                        if little_text:
                            text = ocr_text
                        else:
                            additional = [line for line in ocr_text.splitlines() if line.strip() not in text]
                            text += "\n" + "\n".join(additional)
                    used_ocr = True
                documents.append(Document(page_content=text, metadata={
                    **metadata, "page": page_number, "ocr": used_ocr, "ocr_confidence": confidence,
                }))
        return documents

    def _load_docx(self, payload: bytes, metadata: dict[str, Any]) -> list[Document]:
        from docx import Document as WordDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        doc = WordDocument(io.BytesIO(payload))
        lines = []
        for element in doc.element.body.iterchildren():
            if isinstance(element, CT_P):
                lines.append(Paragraph(element, doc).text)
            elif isinstance(element, CT_Tbl):
                for row in Table(element, doc).rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        image_hashes: set[str] = set()
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            if len(media) > self.max_ocr_images:
                raise DocumentValidationError("Word 图片数超过 OCR 限制")
            for name in media:
                if Path(name).suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
                    continue
                blob = archive.read(name)
                image_hash = hashlib.sha256(blob).hexdigest()
                if image_hash in image_hashes:
                    continue
                image_hashes.add(image_hash)
                text, _confidence = self._ocr_bytes(blob)
                if text:
                    lines.extend([f"[图片文字 {Path(name).name}]", text])
        return [Document(page_content="\n".join(lines), metadata={
            **metadata, "page": 1, "ocr": bool(image_hashes), "ocr_image_count": len(image_hashes),
        })]

    def _load_pptx(self, payload: bytes, metadata: dict[str, Any]) -> list[Document]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        presentation = Presentation(io.BytesIO(payload))
        if len(presentation.slides) > self.max_pages:
            raise DocumentValidationError("PPTX 页数超过解析限制")
        documents = []
        image_count = 0
        for slide_number, slide in enumerate(presentation.slides, start=1):
            lines: list[str] = []
            used_ocr = False

            def visit(shape: Any) -> None:
                nonlocal image_count, used_ocr
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    for child in sorted(shape.shapes, key=lambda item: (item.top, item.left)):
                        visit(child)
                if shape.has_text_frame:
                    lines.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    if image_count > self.max_ocr_images:
                        raise DocumentValidationError("PPTX 图片数超过 OCR 限制")
                    text, _confidence = self._ocr_bytes(shape.image.blob)
                    lines.append(text)
                    used_ocr = True

            for shape in sorted(slide.shapes, key=lambda item: (item.top, item.left)):
                visit(shape)
            documents.append(Document(page_content="\n".join(lines), metadata={
                **metadata, "page": slide_number, "ocr": used_ocr,
            }))
        return documents


def normalize_faq_query(query: str) -> str:
    """Canonicalization shared by Redis and the durable MySQL FAQ lookup."""
    normalized = unicodedata.normalize("NFKC", str(query)).lower()
    return re.sub(r"[\s\W_]+", "", normalized, flags=re.UNICODE)
