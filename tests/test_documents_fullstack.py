"""Real parser checks; fixtures are generated locally and never use old data."""
from __future__ import annotations

import io
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from cloudcare.documents import DocumentProcessor, DocumentValidationError
from cloudcare.settings import Settings


class DocumentProcessorTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.processor = DocumentProcessor(Settings(parent_chunk_size=120, child_chunk_size=50, chunk_overlap=10))

    def test_langchain_parent_child_metadata_and_deterministic_content_hash(self):
        path = self.directory / "客服流程.md"
        path.write_text("# 工单升级\n" + "客户提交故障工单后，客服核验日志，记录处置步骤并通知值班人员。\n" * 20, encoding="utf-8")
        first = self.processor.ingest_path(path, category="工单管理")
        second = self.processor.ingest_path(path, category="工单管理")
        self.assertGreater(len(first.parents), 1)
        self.assertGreater(len(first.chunks), len(first.parents))
        self.assertEqual(first.sha256, second.sha256)
        self.assertEqual([row["id"] for row in first.chunks], [row["id"] for row in second.chunks])
        parents = {row["id"]: row for row in first.parents}
        for row in first.chunks:
            self.assertIn(row["content"], parents[row["parent_id"]]["content"])
            self.assertEqual(row["parent_content"], parents[row["parent_id"]]["content"])
            self.assertEqual(row["tenant_id"], "demo")
            self.assertFalse(row["synthetic"])
            self.assertEqual(len(row["content_hash"]), 64)
        self.assertIn("RecursiveCharacterTextSplitter", first.metadata["splitter"])

    def test_rejects_extension_disguise_invalid_utf8_and_size(self):
        for suffix, content in ((".pdf", b"plain text"), (".txt", b"%PDF-1.7"),
                                (".docx", b"not office"), (".md", b"\xff\xfe")):
            path = self.directory / ("bad" + suffix)
            path.write_bytes(content)
            with self.subTest(suffix=suffix), self.assertRaises(DocumentValidationError):
                self.processor.ingest_path(path)
        path = self.directory / "large.txt"
        path.write_bytes(b"123456789")
        limited = DocumentProcessor(replace(Settings(), max_upload_bytes=8))
        with self.assertRaises(DocumentValidationError):
            limited.ingest_path(path)

    def test_docx_paragraphs_and_table_are_loaded(self):
        from docx import Document
        document = Document()
        document.add_paragraph("CloudCare 工单升级规则")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "故障等级"
        table.cell(0, 1).text = "一级故障由值班人员接管"
        path = self.directory / "support.docx"
        document.save(path)
        loaded = self.processor.ingest_path(path)
        text = "\n".join(d.page_content for d in loaded.documents)
        self.assertIn("工单升级规则", text)
        self.assertIn("一级故障", text)

    def test_pptx_slide_and_table_are_loaded(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text = "CloudCare 退款验收"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(8), Inches(1)).table
        table.cell(0, 0).text = "步骤"
        table.cell(0, 1).text = "核验账单状态"
        path = self.directory / "support.pptx"
        presentation.save(path)
        loaded = self.processor.ingest_path(path)
        text = "\n".join(d.page_content for d in loaded.documents)
        self.assertIn("退款验收", text)
        self.assertIn("核验账单", text)

    def test_text_pdf_preserves_page_provenance(self):
        import fitz
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "CloudCare support incident escalation steps and acceptance rules")
        path = self.directory / "text.pdf"
        pdf.save(path)
        pdf.close()
        loaded = self.processor.ingest_path(path)
        self.assertIn("incident escalation", loaded.documents[0].page_content)
        self.assertEqual(loaded.documents[0].metadata["page"], 1)
        self.assertFalse(loaded.documents[0].metadata["ocr"])

    def test_image_and_scanned_pdf_use_actual_rapidocr(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        font_paths = [Path("C:/Windows/Fonts/msyh.ttc"),
                      Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
        font_path = next((path for path in font_paths if path.exists()), None)
        if font_path is None:
            self.skipTest("OCR fixture requires an installed TrueType font")
        image = Image.new("RGB", (1500, 500), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(str(font_path), 58)
        draw.text((70, 100), "CloudCare support ticket", font=font, fill="black")
        draw.text((70, 240), "Refund within 7 days", font=font, fill="black")
        png = self.directory / "scan.png"
        image.save(png)
        loaded = self.processor.ingest_path(png)
        extracted = loaded.documents[0].page_content.lower().replace(" ", "")
        self.assertIn("refund", extracted)
        self.assertIn("7", extracted)
        self.assertTrue(loaded.documents[0].metadata["ocr"])
        pdf = fitz.open()
        page = pdf.new_page(width=750, height=250)
        page.insert_image(page.rect, filename=str(png))
        path = self.directory / "scan.pdf"
        pdf.save(path)
        pdf.close()
        scanned = self.processor.ingest_path(path)
        self.assertIn("refund", scanned.documents[0].page_content.lower())
        self.assertEqual(scanned.metadata["ocr_pages"], 1)
        self.assertEqual(scanned.metadata["ocr_engine"], "RapidOCR ONNX Runtime")
        mixed_pdf = fitz.open()
        page = mixed_pdf.new_page(width=750, height=330)
        page.insert_text((30, 30), "CloudCare searchable header with scanned customer support body")
        page.insert_image(fitz.Rect(0, 70, 750, 320), filename=str(png))
        mixed_path = self.directory / "mixed.pdf"
        mixed_pdf.save(mixed_path)
        mixed_pdf.close()
        mixed = self.processor.ingest_path(mixed_path)
        self.assertIn("refund", mixed.documents[0].page_content.lower())
        self.assertIn("searchable header", mixed.documents[0].page_content.lower())
        self.assertTrue(mixed.documents[0].metadata["ocr"])


if __name__ == "__main__":
    unittest.main()
